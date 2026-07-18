"""Portfolio slide builder — reusable library.

Refactored from build_slide.py so the same logic can be driven from the
CLI, the Cowork bridge MCP server, or an Agent Skill. All entry points
accept a ticker list and lookback window instead of hardcoding them.
"""

import datetime
import io
import itertools
import time

import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

DEFAULT_TICKERS = ["DTIL", "PRME", "BEAM", "EDIT", "CRSP", "NTLA"]

# Fixed palette for the default portfolio; extra tickers cycle through it.
KNOWN_COLORS = {
    "DTIL": "#1F6FEB",   # blue  (primary focus)
    "PRME": "#8957E5",
    "BEAM": "#3FB950",
    "EDIT": "#F78166",
    "CRSP": "#D29922",
    "NTLA": "#58A6FF",
}
PALETTE_CYCLE = ["#1F6FEB", "#8957E5", "#3FB950", "#F78166", "#D29922", "#58A6FF"]

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
    "Accept": "application/json",
}


def _colors_for(tickers):
    cycle = itertools.cycle(PALETTE_CYCLE)
    return {t: KNOWN_COLORS.get(t) or next(cycle) for t in tickers}


def fetch_history(ticker, days):
    end_ts = int(time.time())
    start_ts = end_ts - ((days + 2) * 24 * 3600)  # pad to ensure full window
    url = (
        f"https://query1.finance.yahoo.com/v8/finance/chart/{ticker}"
        f"?interval=1d&period1={start_ts}&period2={end_ts}"
    )
    r = requests.get(url, headers=HEADERS, timeout=15)
    r.raise_for_status()
    data = r.json()
    result = data["chart"]["result"][0]
    timestamps = result["timestamp"]
    closes = result["indicators"]["quote"][0]["close"]
    dates = [datetime.datetime.utcfromtimestamp(t).strftime("%b %d") for t in timestamps]
    return dates, closes


def fetch_performance(tickers=None, days=30, log=print):
    """Fetch price history and compute per-ticker performance.

    Returns dict with keys: history, perf, current_prices, errors.
    """
    tickers = list(tickers or DEFAULT_TICKERS)
    history, perf, current_prices, errors = {}, {}, {}, {}

    for ticker in tickers:
        try:
            dates, closes = fetch_history(ticker, days)
            valid = [(d, c) for d, c in zip(dates, closes) if c is not None]
            if len(valid) >= 2:
                dd, cc = zip(*valid)
                history[ticker] = (list(dd), list(cc))
                pct = ((cc[-1] - cc[0]) / cc[0]) * 100
                perf[ticker] = round(pct, 2)
                current_prices[ticker] = cc[-1]
                log(f"  {ticker}: {pct:+.2f}%  (${cc[-1]:.2f})")
            else:
                perf[ticker] = None
                history[ticker] = None
                errors[ticker] = "not enough valid price points"
            time.sleep(0.3)
        except Exception as e:
            log(f"  {ticker}: ERROR — {e}")
            perf[ticker] = None
            history[ticker] = None
            errors[ticker] = str(e)

    return {
        "tickers": tickers,
        "days": days,
        "history": history,
        "perf": perf,
        "current_prices": current_prices,
        "errors": errors,
    }


def _build_line_chart(data, focus_ticker):
    tickers = data["tickers"]
    history = data["history"]
    colors = _colors_for(tickers)

    fig, ax = plt.subplots(1, 1, figsize=(11, 5.2), facecolor="#0D1117")
    ax.set_facecolor("#0D1117")

    for ticker in tickers:
        if history.get(ticker) is None:
            continue
        dd, cc = history[ticker]
        base = cc[0]
        norm = [(c / base - 1) * 100 for c in cc]
        is_focus = ticker == focus_ticker
        ax.plot(range(len(dd)), norm, color=colors[ticker],
                linewidth=2.5 if is_focus else 1.5,
                label=ticker, zorder=3 if is_focus else 2)

    ax.axhline(0, color="#444C56", linewidth=0.8, linestyle="--", zorder=1)
    ax.set_xticks([])
    ax.set_ylabel("Return (%)", color="#8B949E", fontsize=9)
    ax.tick_params(colors="#8B949E", labelsize=8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["bottom"].set_color("#30363D")
    ax.spines["left"].set_color("#30363D")
    ax.yaxis.label.set_color("#8B949E")
    plt.setp(ax.get_yticklabels(), color="#8B949E")

    legend = ax.legend(
        loc="upper left", frameon=True, fancybox=False,
        framealpha=0.3, edgecolor="#30363D",
        labelcolor="white", fontsize=8.5, ncol=3,
    )
    legend.get_frame().set_facecolor("#161B22")

    fig.tight_layout(pad=1.2)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor="#0D1117")
    plt.close(fig)
    buf.seek(0)
    return buf


def _build_bar_chart(data):
    perf = data["perf"]
    days = data["days"]
    sorted_tickers = sorted([t for t in data["tickers"] if perf.get(t) is not None],
                            key=lambda t: perf[t])

    fig, ax = plt.subplots(figsize=(5.5, 3.5), facecolor="#0D1117")
    ax.set_facecolor("#0D1117")

    bar_colors = [("#3FB950" if perf[t] >= 0 else "#F78166") for t in sorted_tickers]
    bars = ax.barh(sorted_tickers, [perf[t] for t in sorted_tickers],
                   color=bar_colors, height=0.55, zorder=3)

    for bar, ticker in zip(bars, sorted_tickers):
        val = perf[ticker]
        xpos = bar.get_width() + (0.3 if val >= 0 else -0.3)
        ha = "left" if val >= 0 else "right"
        ax.text(xpos, bar.get_y() + bar.get_height() / 2,
                f"{val:+.1f}%", va="center", ha=ha,
                color="white", fontsize=9, fontweight="bold")

    ax.axvline(0, color="#444C56", linewidth=0.8)
    ax.set_xlabel(f"{_window_label(days)} Return (%)", color="#8B949E", fontsize=8)
    ax.tick_params(colors="white", labelsize=9)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["bottom"].set_color("#30363D")
    ax.spines["left"].set_color("#30363D")
    plt.setp(ax.get_xticklabels(), color="#8B949E")
    fig.tight_layout(pad=1.0)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor="#0D1117")
    plt.close(fig)
    buf.seek(0)
    return buf


def _window_label(days):
    if days % 30 == 0:
        months = days // 30
        return "1-Month" if months == 1 else f"{months}-Month"
    return f"{days}-Day"


def build_deck(data, out_path, title=None):
    """Render the performance deck for already-fetched `data` to out_path."""
    tickers = data["tickers"]
    perf = data["perf"]
    current_prices = data["current_prices"]
    days = data["days"]
    focus_ticker = tickers[0]

    chart_buf = _build_line_chart(data, focus_ticker)
    bar_buf = _build_bar_chart(data)

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string("0D1117")

    def add_text(text, left, top, width, height, font_size, bold=False,
                 color="FFFFFF", align=PP_ALIGN.LEFT):
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                         Inches(width), Inches(height))
        tf = txBox.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        return txBox

    add_text(title or f"Biotech Portfolio — {_window_label(days)} Performance",
             0.3, 0.15, 9, 0.55, 22, bold=True, color="E6EDF3")

    today = datetime.date.today()
    window_ago = today - datetime.timedelta(days=days)
    subtitle = (f"{window_ago.strftime('%b %d')} – {today.strftime('%b %d, %Y')}"
                f"  |  {' · '.join(tickers)}")
    add_text(subtitle, 0.3, 0.72, 12, 0.35, 9, color="8B949E")

    line = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.RECTANGLE
        Inches(0.3), Inches(1.08), Inches(12.73), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0x21, 0x26, 0x2D)
    line.line.fill.background()

    slide.shapes.add_picture(chart_buf, Inches(0.3), Inches(1.15),
                             Inches(7.8), Inches(4.0))
    slide.shapes.add_picture(bar_buf, Inches(8.25), Inches(1.15),
                             Inches(4.8), Inches(3.8))

    card_w = 2.0
    card_h = 0.75
    start_x = 0.3
    card_y = 5.6

    for i, ticker in enumerate(tickers[:6]):
        x = start_x + i * (card_w + 0.15)
        card = slide.shapes.add_shape(
            1, Inches(x), Inches(card_y), Inches(card_w), Inches(card_h)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0x16, 0x1B, 0x22)
        card.line.color.rgb = RGBColor(0x30, 0x36, 0x3D)

        add_text(ticker, x + 0.1, card_y + 0.04, card_w - 0.2, 0.3,
                 10, bold=True, color="E6EDF3")

        if perf.get(ticker) is not None:
            pct = perf[ticker]
            price = current_prices.get(ticker, 0)
            pct_color = "3FB950" if pct >= 0 else "F78166"
            sign = "+" if pct >= 0 else ""
            add_text(f"{sign}{pct:.2f}%",
                     x + 0.1, card_y + 0.33, card_w - 0.2, 0.28,
                     13, bold=True, color=pct_color)
            add_text(f"${price:.2f}",
                     x + 1.1, card_y + 0.37, 0.8, 0.28,
                     9, bold=False, color="8B949E", align=PP_ALIGN.RIGHT)
        else:
            add_text("N/A", x + 0.1, card_y + 0.33, card_w - 0.2, 0.28,
                     11, bold=True, color="8B949E")

    add_text("Source: Yahoo Finance  |  For internal use only — not for distribution",
             0.3, 7.1, 12.73, 0.3, 7.5, color="444C56", align=PP_ALIGN.CENTER)

    prs.save(out_path)
    return out_path


def build_portfolio_slide(tickers=None, days=30, out_path=None, title=None, log=print):
    """Fetch data and render the deck in one call. Returns (out_path, data)."""
    import os

    data = fetch_performance(tickers, days, log=log)
    if out_path is None:
        repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        out_path = os.path.join(
            repo_root, f"Biotech_Portfolio_{_window_label(days).replace('-', '')}.pptx"
        )
    build_deck(data, out_path, title=title)
    return out_path, data
